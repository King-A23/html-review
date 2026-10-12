"""Create fictional, openly distributable DOCX/PDF review examples.
Run with Python having python-docx and reportlab installed.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'examples/documents'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Subtitle']:
    style = doc.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0,0,0)
    for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(10)
doc.add_paragraph('Harbor pilot review', 'Title')
doc.add_paragraph('Fictional project report for document review practice', 'Subtitle')
doc.add_paragraph('We recommend a second, smaller pilot before a wider launch. The four-week trial improved task completion, but the report needs clearer metric definitions and a named owner for the next step.')
doc.add_heading('Pilot results', 1)
doc.add_paragraph('The pilot included 40 fictional participants across two groups. Completion rose from 68% to 84%. These figures are illustrative and do not describe a real organization.')
table = doc.add_table(rows=1, cols=3)
for c,t in zip(table.rows[0].cells, ['Measure','Before','After']): c.text=t
for values in [('Completion','68%','84%'),('Time per task','12 min','9 min'),('Support requests','18','11')]:
    for c,t in zip(table.add_row().cells, values): c.text=t
for i,row in enumerate(table.rows):
    for cell in row.cells:
        tcPr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        tcPr.append(borders)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF6' if i==0 else ('F5F7FA' if i%2==0 else 'FFFFFF'));tcPr.append(sh)
        for p in cell.paragraphs: p.paragraph_format.space_after=Pt(8);p.paragraph_format.space_before=Pt(8)
doc.add_paragraph().paragraph_format.space_after = Pt(2)
doc.add_paragraph('Review note: define how completion was measured before comparing groups.')
doc.add_heading('What we learned',1)
doc.add_paragraph('Shorter instructions helped participants finish the first task. The second group still asked for examples, suggesting that the next pilot should test a guided introduction.')
doc.add_page_break()
doc.add_heading('Next pilot plan',1)
doc.add_paragraph('Run a two-week pilot with 20 new fictional participants. Keep the same task and measure completion without assistance, then compare the results with the first pilot.')
doc.add_heading('Decisions to review',1)
doc.add_paragraph('Assign an owner for the next pilot and confirm the start date. The current draft does not identify either, so the schedule cannot yet be approved.')
doc.add_paragraph('Publish the metric definition with the final report. Include the sample size and the number of completed tasks so the percentage can be checked.')
doc.add_heading('Release checklist',1)
for text in ['Define the completion metric and denominator.','Name the pilot owner and confirm the start date.','Keep the report, source files and review comments together.']:
    doc.add_paragraph(text,'List Bullet')
doc.add_paragraph('All names, numbers and plans in this example are fictional. No external links or personal information are included.')
doc.core_properties.author='HTML Review examples';doc.core_properties.title='Harbor pilot review';doc.core_properties.subject='Fictional document review example'
doc.save(OUT/'review.docx')
c=canvas.Canvas(str(OUT/'review.pdf'),pagesize=A4,invariant=1)
c.setTitle('Harbor pilot review');c.setAuthor('HTML Review examples')
w,h=A4
NAVY=HexColor('#17324D');GREY=HexColor('#526173')
def text(y,t,size=11,color=NAVY):
    c.setFillColor(color);c.setFont('Helvetica',size);c.drawString(54,y,t)
def base(n,title):
    text(790,'FICTIONAL PROJECT / REVIEW EXAMPLE',9,GREY)
    text(750,title,25)
    c.setStrokeColor(HexColor('#D9E2EC'));c.line(54,730,w-54,730)
    text(35,f'Harbor pilot  |  Illustrative data only                                      {n} / 2',9,GREY)
base(1,'Harbor pilot review')
for y,t in [(697,'Recommendation: run a smaller second pilot before a wider launch.'),(674,'The four-week trial improved completion, but the metric definition'),(656,'and the next-step owner still need review.')]:text(y,t)
text(610,'Pilot results',16)
rows=[('Measure','Before','After'),('Completion','68%','84%'),('Time per task','12 min','9 min'),('Support requests','18','11')]
for i,row in enumerate(rows):
    y=567-i*36;c.setFillColor(HexColor('#E8EFF6' if i==0 else '#F5F7FA'));c.rect(54,y-12,w-108,35,fill=1,stroke=0)
    for x,t in zip([66,310,420],row):
        c.setFillColor(NAVY);c.setFont('Helvetica',11);c.drawString(x,y,t)
text(392,'Completion improved by 16 percentage points.',12)
text(370,'Review the definition and denominator before approving this claim.',10,GREY)
text(320,'Task completion',16)
for y,label,value in [(277,'Before',68),(235,'After',84)]:
    text(y,label,10);c.setFillColor(HexColor('#7192B3' if value==68 else '#235E87'));c.rect(120,y-4,value*3.7,18,fill=1,stroke=0);c.setFillColor(NAVY);c.setFont('Helvetica',10);c.drawString(440,y,f'{value}%')
text(175,'Sample: 40 fictional participants across two groups.',10,GREY)
text(150,'All data is invented for review practice. No personal information.',10,GREY)
c.showPage();base(2,'Next pilot plan')
text(690,'Scope',16)
for y,t in [(660,'Run a two-week pilot with 20 new fictional participants.'),(640,'Use the same task and measure completion without assistance.')]:text(y,t)
text(580,'Decisions to review',16)
for y,t in [(548,'1. Define completion and publish the denominator.'),(518,'2. Assign an owner and confirm the pilot start date.'),(488,'3. Retain source files with the review comments.')]:text(y,t)
text(410,'Open question',16)
text(380,'The current draft has no named owner or confirmed start date.')
text(359,'The schedule therefore remains a proposal, pending review.')
text(270,'Review practice',16)
text(240,'Select text to request a clearer definition. Mark the chart region')
text(220,'to request a label, then export the saved review comments.')
c.save()
print('Created examples/documents/review.docx and review.pdf')
