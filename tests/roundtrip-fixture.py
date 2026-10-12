#!/usr/bin/env python3
"""Rebuild a fictional, local DOCX review fixture using bundled Python.

Run the documents skill authoring marker before the first authoring invocation.
This fixture contains intentional editorial defects for a real review roundtrip.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'iteration-review'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}

def chart(path, color='#178C92'):
    im = Image.new('RGB', (1440, 760), 'white')
    d = ImageDraw.Draw(im)
    font = '/System/Library/Fonts/Supplemental/Arial.ttf'
    label = ImageFont.truetype(font, 34)
    small = ImageFont.truetype(font, 28)
    bold = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf', 38)
    d.text((95, 35), 'Quarterly revenue', fill='#202A30', font=bold)
    d.text((95, 91), '2026 Q1-Q3 | CNY 100 million', fill='#607078', font=small)
    left, right, top, bottom = 150, 1350, 175, 620
    for v in (0, 1, 2, 3):
        y = bottom - v / 3.4 * (bottom-top)
        d.line((left, y, right, y), fill='#E1E7EA', width=2)
        d.text((88, y-18), str(v), fill='#607078', font=small)
    for x, title, value in [(340, 'Q1', 2.4), (740, 'Q2', 2.7), (1140, 'Q3', 3.0)]:
        y = bottom-value/3.4*(bottom-top)
        d.rectangle((x-100, y, x+100, bottom), fill=color)
        d.text((x, y-55), f'{value:.1f}', fill='#202A30', font=bold, anchor='mt')
        d.text((x, bottom+25), title, fill='#202A30', font=label, anchor='mt')
    d.text((95, 705), 'Source: Chengchuan operating ledger (fictional)', fill='#607078', font=small)
    im.save(path)

def fontset(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'PingFang SC')

def para(doc, text='', style=None):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.widow_control = True
    return p

def footnote_patch(path):
    with ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    body = etree.fromstring(files['word/document.xml'])
    target = next(p for p in body.xpath('//w:p', namespaces=NS) if '比较口径以累计同比为准' in ''.join(p.xpath('.//w:t/text()', namespaces=NS)))
    r = etree.SubElement(target, qn('w:r'))
    pr = etree.SubElement(r, qn('w:rPr'))
    etree.SubElement(pr, qn('w:rStyle')).set(qn('w:val'), 'FootnoteReference')
    etree.SubElement(r, qn('w:footnoteReference')).set(qn('w:id'), '1')
    notes = etree.Element(qn('w:footnotes'), nsmap={'w': W})
    for nid, typ, tag in [('-1', 'separator', 'separator'), ('0', 'continuationSeparator', 'continuationSeparator')]:
        note = etree.SubElement(notes, qn('w:footnote')); note.set(qn('w:id'), nid); note.set(qn('w:type'), typ)
        p = etree.SubElement(note, qn('w:p')); rr = etree.SubElement(p, qn('w:r')); etree.SubElement(rr, qn('w:'+tag))
    note = etree.SubElement(notes, qn('w:footnote')); note.set(qn('w:id'), '1')
    p = etree.SubElement(note, qn('w:p'))
    rr = etree.SubElement(p, qn('w:r')); etree.SubElement(rr, qn('w:footnoteRef'))
    rr = etree.SubElement(p, qn('w:r')); etree.SubElement(rr, qn('w:t')).text = ' 本报告所涉公司、客户与数据均为虚构。同比基期收入为6.75亿元，8.10 ÷ 6.75 − 1 = 20%。'
    relpath = 'word/_rels/document.xml.rels'
    rels = etree.fromstring(files[relpath])
    rel = etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
    rel.set('Id', 'rIdFixtureFootnote'); rel.set('Type', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'); rel.set('Target', 'footnotes.xml')
    ct = etree.fromstring(files['[Content_Types].xml'])
    o = etree.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override')
    o.set('PartName', '/word/footnotes.xml'); o.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    for name, xml in [('word/document.xml', body), ('word/footnotes.xml', notes), (relpath, rels), ('[Content_Types].xml', ct)]:
        files[name] = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, data in files.items(): z.writestr(name, data)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out-dir', type=Path, default=OUT)
    args = parser.parse_args()
    out = args.out_dir.resolve(); out.mkdir(parents=True, exist_ok=True)
    image = out/'季度营收.png'; chart(image)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    for name, size, bold in [('Normal', 11, False), ('Title', 24, True), ('Subtitle', 12, False), ('Heading 1', 16, True), ('Heading 2', 12, True), ('Caption', 10, False), ('List Number', 11, False)]:
        fontset(doc.styles[name], size, bold)
    for border in list(doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(9); normal.line_spacing = 1.3
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(16)
        doc.styles[name].paragraph_format.space_after = Pt(8)
    doc.core_properties.title = '澄川科技2026年前三季度经营报告'
    doc.core_properties.author = '澄川科技经营分析组'
    para(doc, '澄川科技2026年前三季度经营报告', 'Title')
    para(doc, '管理层审阅稿  2026年10月8日', 'Subtitle')
    para(doc, '本报告面向管理层，评估前三季度的收入、利润与现金回收，并提出第四季度的执行重点。澄川科技是一家虚构的企业服务公司；文内数据均为演示数据，不对应任何真实公司或投资标的。')
    para(doc, '一 经营摘要', 'Heading 1')
    p = para(doc)
    pieces = [('2026年前三季度，公司累计实现营业收入', False, False), ('8.1亿元', True, False), ('，同比增长', False, False), ('20%', True, False), ('；经营活动现金流净额为', False, False), ('0.90亿元', True, False), ('，客户回款节奏改善。', False, False), ('Revenue growth', False, True), ('主要来自存量客户续约与实施服务扩张，净利润累计为0.81亿元，利润率保持10%。团队持续跟进签约、交付、验收和回款链路，并按月复盘项目排期、交付效率、客户使用率和合同执行情况。整体经营表现符合内部计划，后续仍需关注大型项目验收时间以及期末应收账款回收。', False, False)]
    for text, bold, italic in pieces:
        run = p.add_run(text); run.bold = bold; run.italic = italic
    para(doc, '现金回收仍需持续关注。')
    para(doc, '二 业务进展', 'Heading 1')
    para(doc, '客户续约与交付', 'Heading 2')
    para(doc, '核心客户续约率为92%，实施项目按期验收率为88%。销售团队以续约与扩容为主，交付团队将资源优先投入已经签署合同的项目，以缩短从订单到收入确认的周期。')
    para(doc, '免责声明 临时重复无效文字 本段为排版测试占位，请勿据此作出任何经营判断。本段为排版测试占位，请勿据此作出任何经营判断。')
    doc.add_page_break()
    para(doc, '三 财务表现', 'Heading 1')
    para(doc, '前三季度收入逐季增加，净利润与收入同步增长。以下金额均以亿元计，比较口径以累计同比为准。')
    para(doc, '表 1 季度财务表现', 'Caption')
    table = doc.add_table(rows=6, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for c, width in zip(table.columns, [4.5, 4.0, 4.0, 4.0]): c.width = Cm(width)
    table.cell(0,0).merge(table.cell(1,0)).text = '期间'
    table.cell(0,1).merge(table.cell(0,2)).text = '经营成果'
    table.cell(0,3).merge(table.cell(1,3)).text = '净利润率'
    table.cell(1,1).text = '营业收入'; table.cell(1,2).text = '净利润'
    data = [('Q1', '2.40', '0.24', '10%'), ('Q2', '2.70', '0.27', '10%'), ('Q3', '3.00', '0.30', '10%'), ('前三季度合计', '8.10', '0.81', '10%')]
    for i, row in enumerate(data, 2):
        for j, text in enumerate(row): table.cell(i,j).text = text
    seen=set()
    for i,row in enumerate(table.rows):
        if i<2:
            header=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(header)
        for cell in row.cells:
            if cell._tc in seen: continue
            seen.add(cell._tc)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                item=OxmlElement('w:'+side); item.set(qn('w:val'),'single'); item.set(qn('w:sz'),'6'); item.set(qn('w:color'),'D9D9D9'); borders.append(item)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                item=OxmlElement('w:'+side); item.set(qn('w:w'),'130'); item.set(qn('w:type'),'dxa'); margins.append(item)
            pr.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCEAF1' if i<2 else ('F5F7F8' if i%2 else 'FFFFFF')); pr.append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
                for run in p.runs: run.bold = i<2 or i==5; run.font.size=Pt(10.5)
    source = para(doc, '来源：澄川科技经营台账，2026年9月30日。金额保留两位小数，表内为虚构数据。', 'Caption')
    source.paragraph_format.space_before = Pt(6)
    para(doc, '现金流与应收账款', 'Heading 2')
    para(doc, '经营活动现金流净额0.90亿元，高于净利润0.81亿元，主要受存量合同回款影响。期末应收账款为1.12亿元，账龄超过90天的部分占18%。财务团队应将逾期催收纳入每周跟踪，并核对客户验收文件。')
    para(doc, '现金回收仍需持续关注。')
    para(doc, '经营约束', 'Heading 2')
    para(doc, '第三季度新增实施项目集中在月末验收，第四季度收入仍取决于客户签收。公司应保持销售与交付排期一致，避免用未完成验收的订单替代已确认收入。')
    doc.add_page_break()
    para(doc, '四 收入趋势与执行重点', 'Heading 1')
    para(doc, '季度营业收入从2.40亿元增至3.00亿元。图中比较同一年度的季度水平，不表示同比增速；管理层应结合回款与交付能力判断增长的持续性。')
    p = para(doc); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(image), width=Cm(16.5))
    picture._inline.docPr.set('descr', '蓝绿色柱状图，Q1收入2.4亿元，Q2收入2.7亿元，Q3收入3.0亿元，图片尺寸1440乘760像素')
    p.paragraph_format.keep_with_next = True
    para(doc, '图 1 2026年前三季度营业收入', 'Caption')
    para(doc, '来源：澄川科技经营台账。图表数据与表1一致，金额单位为亿元。', 'Caption')
    para(doc, '第四季度执行安排', 'Heading 2')
    for text in ['经营分析组于每月第五个工作日前更新收入与现金流预测，分别列示已验收项目和待验收项目。', '客户成功团队跟踪重点客户续约风险，采用Customer Success复盘记录解释使用率与回款变化。', '财务团队每周复核逾期应收账款，交付团队同步补齐验收依据，并在月度经营会上报告处理进度。']:
        para(doc, text, 'List Number')
    para(doc, '管理层在10月经营会上确认资源分配与催收负责人，11月按实际验收和回款情况调整预测。以上安排为虚构情境中的管理建议。')
    path = out/'澄川经营报告.docx'; doc.save(path); footnote_patch(path)
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith('.xml') or name.endswith('.rels'): etree.fromstring(z.read(name))
        assert not any(b'TargetMode="External"' in z.read(n) for n in z.namelist() if n.endswith('.rels'))
        assert Image.open(BytesIO(z.read('word/media/image1.png'))).size == (1440,760)
    print(path)

if __name__ == '__main__': main()
